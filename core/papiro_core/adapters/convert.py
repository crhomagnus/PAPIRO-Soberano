# -*- coding: utf-8 -*-
"""Conversao a partir do PDF (RF-005, RF-602 parcial, RF-603), formularios (RF-701..RF-706)
e seguranca basica (RF-805 criptografia, RF-809 sanitizacao)."""
from __future__ import annotations
import csv, json, pathlib, secrets, string
from collections import Counter
from xml.sax.saxutils import escape as xml_escape
import fitz, pikepdf
from .. import fidelidade as FID
from ..erros import PapiroErro
from .inspect import abrir, abrir_pike, triagem_risco

N = pikepdf.Name


# ---------------- conversao ----------------
def para_texto(entrada: pathlib.Path, out: pathlib.Path) -> dict:
    with abrir(entrada) as doc:
        txt = "\n".join(f"--- pagina {i + 1} ---\n{p.get_text()}" for i, p in enumerate(doc))
        n = doc.page_count
    out.write_text(txt, encoding="utf-8")
    return {"paginas": n}


def para_markdown(entrada: pathlib.Path, out: pathlib.Path) -> dict:
    try:
        from pymupdf4llm import to_markdown
        out.write_text(to_markdown(str(entrada)), encoding="utf-8")
        return {"via": "pymupdf4llm"}
    except ImportError:
        with abrir(entrada) as doc:
            out.write_text("\n\n".join(f"<!-- pagina {i + 1} -->\n\n{p.get_text()}" for i, p in enumerate(doc)),
                           encoding="utf-8")
        return {"via": "pymupdf", "fallback_from": "pymupdf4llm"}


def para_json(entrada: pathlib.Path, out: pathlib.Path) -> dict:
    """Estrutura por pagina: blocos de texto com caixa, fonte e corpo (ordem de leitura do PyMuPDF)."""
    paginas = []
    with abrir(entrada) as doc:
        for page in doc:
            blocos = []
            for b in page.get_text("dict", sort=True)["blocks"]:
                if b.get("type") != 0:
                    blocos.append({"tipo": "imagem", "bbox": [round(v, 1) for v in b["bbox"]]})
                    continue
                spans = [s for ln in b["lines"] for s in ln["spans"]]
                blocos.append({"tipo": "texto", "bbox": [round(v, 1) for v in b["bbox"]],
                               "texto": " ".join("".join(s["text"] for s in ln["spans"]) for ln in b["lines"]).strip(),
                               "fonte": spans[0]["font"] if spans else "", "corpo": round(spans[0]["size"], 1) if spans else 0})
            paginas.append({"pagina": page.number + 1, "largura": page.rect.width, "altura": page.rect.height, "blocos": blocos})
    out.write_text(json.dumps({"paginas": paginas}, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"paginas": len(paginas)}


def para_svg(entrada: pathlib.Path, out_dir: pathlib.Path) -> list[pathlib.Path]:
    outs = []
    with abrir(entrada) as doc:
        for page in doc:
            p = out_dir / f"pag{page.number + 1:03d}.svg"
            p.write_text(page.get_svg_image(text_as_path=False), encoding="utf-8")
            outs.append(p)
    return outs


def para_png(entrada: pathlib.Path, out_dir: pathlib.Path, dpi: int = 110, prancha: bool = False) -> list[pathlib.Path]:
    """RF-005: PNG por pagina com pypdfium2 (fallback PyMuPDF) + prancha de contato opcional."""
    out_dir.mkdir(parents=True, exist_ok=True)
    outs: list[pathlib.Path] = []
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(entrada))
        try:
            for i in range(len(pdf)):
                p = out_dir / f"pag{i + 1:03d}.png"
                pdf[i].render(scale=dpi / 72).to_pil().save(p)
                outs.append(p)
        finally:
            pdf.close()
    except ImportError:
        with abrir(entrada) as doc:
            for i, page in enumerate(doc):
                p = out_dir / f"pag{i + 1:03d}.png"
                page.get_pixmap(dpi=dpi).save(p)
                outs.append(p)
    if prancha and outs:
        from PIL import Image
        miniaturas = []
        for p in outs:
            im = Image.open(p)
            im.thumbnail((240, 340))
            miniaturas.append(im.convert("RGB"))
        cols = min(4, len(miniaturas))
        lins = -(-len(miniaturas) // cols)
        folha = Image.new("RGB", (cols * 250, lins * 350), "white")
        for k, im in enumerate(miniaturas):
            folha.paste(im, ((k % cols) * 250 + 5, (k // cols) * 350 + 5))
        pr = out_dir / "prancha.png"
        folha.save(pr)
        outs.append(pr)
    return outs


def tabelas(entrada: pathlib.Path, out: pathlib.Path) -> dict:
    """RF-603: pdfplumber (primario) x PyMuPDF find_tables (conferencia); divergencia celula a celula."""
    import pdfplumber
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "tabelas"
    wd = wb.create_sheet("divergencias")
    wd.append(["pagina", "tabela", "linha", "coluna", "pdfplumber", "pymupdf"])
    total, divergencias = 0, 0
    with pdfplumber.open(entrada) as pdf, abrir(entrada) as doc:
        for pno, page in enumerate(pdf.pages, start=1):
            t_plumber = page.extract_tables() or []
            try:
                t_fitz = [t.extract() for t in doc[pno - 1].find_tables().tables]
            except Exception:
                t_fitz = []
            for ti, t in enumerate(t_plumber, start=1):
                ws.append([f"pagina {pno} - tabela {ti}"])
                outra = t_fitz[ti - 1] if ti - 1 < len(t_fitz) else None
                for li, row in enumerate(t, start=1):
                    ws.append(row)
                    total += 1
                    for ci, cel in enumerate(row, start=1):
                        ref = None
                        if outra is not None and li - 1 < len(outra) and ci - 1 < len(outra[li - 1]):
                            ref = outra[li - 1][ci - 1]
                        norm = lambda v: " ".join(str(v or "").split())
                        if outra is None or norm(cel) != norm(ref):
                            wd.append([pno, ti, li, ci, cel, ref])
                            divergencias += 1
    if out.suffix.lower() == ".csv":
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(row)
        xlsx = out.with_suffix(".divergencias.xlsx")
        wb.save(xlsx)
    else:
        wb.save(out)
    return {"linhas": total, "divergencias": divergencias}


# ---------------- formularios ----------------
def listar_campos(entrada: pathlib.Path) -> list[dict]:
    out = []
    with abrir(entrada) as doc:
        for page in doc:
            for w in page.widgets() or []:
                out.append({"pagina": page.number + 1, "nome": w.field_name, "tipo": w.field_type_string,
                            "valor": w.field_value, "rect": [round(v, 1) for v in w.rect],
                            "opcoes": list(w.choice_values or []), "somente_leitura": bool(w.field_flags & 1)})
    return out


def detectar_xfa(entrada: pathlib.Path) -> bool:
    with abrir_pike(entrada) as pdf:
        acro = pdf.Root.get(N.AcroForm)
        return acro is not None and N.XFA in acro


def preencher(entrada: pathlib.Path, out: pathlib.Path, dados: dict) -> dict:
    """RF-703: preenche por nome, regenera aparencia; campos desconhecidos sao informados."""
    avisos = []
    if detectar_xfa(entrada):
        avisos.append("formulario XFA: somente a camada AcroForm foi preenchida (RF-706)")
    usados = set()
    with abrir(entrada) as doc:
        for page in doc:
            for w in page.widgets() or []:
                if w.field_name in dados:
                    valor = dados[w.field_name]
                    if w.field_type in (fitz.PDF_WIDGET_TYPE_CHECKBOX, fitz.PDF_WIDGET_TYPE_RADIOBUTTON):
                        w.field_value = w.on_state() if valor in (True, "true", "sim", "1", 1, w.on_state()) else "Off"
                    else:
                        w.field_value = str(valor)
                    w.update()
                    usados.add(w.field_name)
        doc.save(out, garbage=3, deflate=True)
    desconhecidos = sorted(set(dados) - usados)
    if desconhecidos:
        avisos.append(f"campos inexistentes no PDF: {desconhecidos}")
    return {"preenchidos": len(usados), "avisos": avisos}


def exportar_campos(entrada: pathlib.Path, out: pathlib.Path) -> dict:
    """RF-704: JSON, CSV, FDF ou XFDF conforme a extensao de `out`."""
    campos = listar_campos(entrada)
    valores = {c["nome"]: ("" if c["valor"] is None else str(c["valor"])) for c in campos}
    suf = out.suffix.lower()
    if suf == ".json":
        out.write_text(json.dumps(valores, ensure_ascii=False, indent=2), encoding="utf-8")
    elif suf == ".csv":
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(valores.keys())
            w.writerow(valores.values())
    elif suf == ".xfdf":
        linhas = ['<?xml version="1.0" encoding="UTF-8"?>',
                  '<xfdf xmlns="http://ns.adobe.com/xfdf/" xml:space="preserve"><fields>']
        linhas += [f'<field name="{xml_escape(k, {chr(34): "&quot;"})}"><value>{xml_escape(v)}</value></field>' for k, v in valores.items()]
        linhas.append("</fields></xfdf>")
        out.write_text("\n".join(linhas), encoding="utf-8")
    elif suf == ".fdf":
        def pdfstr(s: str) -> str:
            return "(" + s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ")"
        def utf16(s: str) -> str:
            return "<FEFF" + s.encode("utf-16-be").hex().upper() + ">"
        corpo = "".join(f"<< /T {pdfstr(k)} /V {utf16(v)} >>" for k, v in valores.items())
        out.write_bytes(("%FDF-1.2\n1 0 obj\n<< /FDF << /Fields [" + corpo + "] >> >>\nendobj\ntrailer\n<< /Root 1 0 R >>\n%%EOF\n").encode("latin-1"))
    else:
        raise PapiroErro("E_ENTRADA", "formato de exportacao: .json, .csv, .fdf ou .xfdf")
    return {"campos": len(valores)}


def importar_xfdf(entrada: pathlib.Path, out: pathlib.Path, xfdf: pathlib.Path) -> dict:
    import xml.etree.ElementTree as ET
    raiz = ET.parse(xfdf).getroot()
    ns = {"x": "http://ns.adobe.com/xfdf/"}
    dados = {f.get("name"): (f.findtext("x:value", default="", namespaces=ns) or f.findtext("value", default=""))
             for f in raiz.iter() if f.tag.endswith("field")}
    return preencher(entrada, out, dados)


def achatar(entrada: pathlib.Path, out: pathlib.Path, anotacoes: bool = True) -> dict:
    """RF-705/RF-207: queima campos (e anotacoes) no conteudo, remove AcroForm e confere aparencia por SSIM."""
    with abrir(entrada) as doc:
        widgets_antes = sum(len(list(p.widgets() or [])) for p in doc)
        doc.bake(annots=anotacoes, widgets=True)
        doc.save(out, garbage=4, deflate=True)
    with pikepdf.open(out, allow_overwriting_input=True) as pdf:
        if N.AcroForm in pdf.Root:
            del pdf.Root[N.AcroForm]
            pdf.save(out)
    with fitz.open(out) as d:
        widgets_depois = sum(len(list(p.widgets() or [])) for p in d)
    if widgets_depois:
        raise PapiroErro("E_MOTOR", f"{widgets_depois} campos continuam editaveis apos o achatamento")
    fid = FID.comparar(entrada, out, dpi=96)
    return {"campos_achatados": widgets_antes, "campos_editaveis_restantes": 0,
            "ssim_pior_bloco": fid["ssim_pior_bloco"], "aparencia_identica": fid["ssim_pior_bloco"] >= 0.99}


# ---------------- seguranca basica ----------------
def gerar_senha(tamanho: int = 20) -> str:
    alfabeto = string.ascii_letters + string.digits + "-_.@#"
    return "".join(secrets.choice(alfabeto) for _ in range(tamanho))


def criptografar(entrada: pathlib.Path, out: pathlib.Path, senha: str = "", senha_dono: str = "",
                 imprimir: bool = True, copiar: bool = False, editar: bool = False) -> dict:
    """RF-805: AES-256 (R6). Permissoes sao respeitadas por convencao do leitor, nao por protecao forte."""
    gerada = not senha
    senha = senha or gerar_senha()
    senha_dono = senha_dono or gerar_senha(32)
    perms = pikepdf.Permissions(print_lowres=imprimir, print_highres=imprimir, extract=copiar,
                                modify_other=editar, modify_annotation=editar, modify_form=editar, modify_assembly=editar)
    with abrir_pike(entrada) as pdf:
        pdf.save(out, encryption=pikepdf.Encryption(owner=senha_dono, user=senha, R=6, allow=perms))
    return {"alg": "AES-256 (R6)", "senha_gerada": gerada, "senha": senha if gerada else None,
            "aviso": "permissoes dependem do leitor respeitar a convencao"}


def decriptar(entrada: pathlib.Path, out: pathlib.Path, senha: str) -> dict:
    if not senha:
        raise PapiroErro("E_SENHA")
    with abrir_pike(entrada, senha) as pdf:
        pdf.save(out)
    return {"ok": True}


ACOES_PERIGOSAS = {N.JavaScript, N.Launch, N.SubmitForm, N.ImportData, N.GoToR, N.GoToE, N.Rendition,
                   N.Sound, N.Movie, N("/RichMediaExecute")}
ANOTACOES_REMOVER = {N.FileAttachment, N.RichMedia, N.Screen, N.Movie, N.Sound, N("/3D")}


def sanitizar(entrada: pathlib.Path, out: pathlib.Path, manter_metadados_basicos: bool = True) -> dict:
    """RF-809/§12.4: remove JavaScript, OpenAction, AA, Launch e demais acoes perigosas, anexos,
    XFA, XMP (inclusive por objeto), miniaturas e dados privados de aplicativos. Exige triagem limpa."""
    rel: Counter = Counter()
    with abrir_pike(entrada) as pdf:
        root = pdf.Root
        for k in (N.OpenAction, N.AA, N.Metadata, N.PieceInfo, N.SpiderInfo):
            if k in root:
                del root[k]
                rel[str(k)] += 1
        nomes = root.get(N.Names)
        if nomes is not None:
            for k in (N.JavaScript, N.EmbeddedFiles):
                if k in nomes:
                    del nomes[k]
                    rel[f"Names{k}"] += 1
        acro = root.get(N.AcroForm)
        if acro is not None and N.XFA in acro:
            del acro[N.XFA]
            rel["/XFA"] += 1
        for page in pdf.pages:
            annots = page.obj.get(N.Annots)
            if annots is not None:
                manter = pikepdf.Array([a for a in annots if a.get(N.Subtype) not in ANOTACOES_REMOVER])
                removidas = len(annots) - len(manter)
                if removidas:
                    page.obj[N.Annots] = manter
                    rel["anotacoes_perigosas"] += removidas
        for obj in pdf.objects:
            if not isinstance(obj, (pikepdf.Dictionary, pikepdf.Stream)):
                continue
            try:
                for k in (N.AA, N.PieceInfo, N.Thumb):
                    if k in obj:
                        del obj[k]
                        rel[str(k)] += 1
                if N.Metadata in obj and obj.get(N.Type) != N.Catalog:
                    del obj[N.Metadata]
                    rel["/Metadata(objeto)"] += 1
                acao = obj.get(N.A)
                if isinstance(acao, pikepdf.Dictionary) and acao.get(N.S) in ACOES_PERIGOSAS:
                    del obj[N.A]
                    rel["acao " + str(acao.get(N.S))] += 1
                if obj.get(N.S) in ACOES_PERIGOSAS:
                    for k in list(obj.keys()):
                        if k != "/Type":
                            del obj[k]
                    obj[N.S] = N.GoTo
                    obj[N.D] = pikepdf.Array([pdf.pages[0].obj, N.Fit])
                    rel["acao neutralizada"] += 1
                if N.Next in obj:
                    del obj[N.Next]
            except Exception:
                continue
        basicos = {}
        if manter_metadados_basicos and pdf.docinfo is not None:
            basicos = {k: pdf.docinfo[k] for k in ("/Title", "/Author", "/Subject", "/Keywords") if k in pdf.docinfo}
        for k in list(pdf.docinfo.keys()):
            del pdf.docinfo[k]
        for k, v in basicos.items():
            pdf.docinfo[k] = v
        pdf.docinfo["/Producer"] = "PAPIRO (sanitizado)"
        if basicos:
            with pdf.open_metadata(set_pikepdf_as_editor=False) as xmp:
                xmp.load_from_docinfo(pdf.docinfo)
        pdf.remove_unreferenced_resources()
        pdf.save(out)
    triagem = triagem_risco(out)
    if triagem["nota_risco"] != "BAIXA":
        raise PapiroErro("E_CONFORMIDADE", f"triagem apos sanitizar continua {triagem['nota_risco']}")
    return {"removidos": dict(rel), "triagem_depois": triagem["nota_risco"]}
