# -*- coding: utf-8 -*-
"""Nivel 3 - criacao RF-301..RF-310.
Markdown -> Typst (primario, texto entra como string literal #"...") ou PyMuPDF Story (fallback);
fontes OFL de fonts/ sempre embutidas; titulo, autor, idioma e produtor gravados (G7)."""
from __future__ import annotations
import csv, io, json, pathlib, re, shutil, tempfile
import fitz, pikepdf
from markdown_it import MarkdownIt
from .. import FONTS, binfinder as BF, config, engines as ENG
from ..erros import PapiroErro
from .edit import metadados

FAMILIA = "Liberation Sans"


# ---------------- Markdown -> Typst ----------------
def _s(texto: str) -> str:
    """String literal Typst: nenhum caractere do usuario vira marcacao."""
    return '#"' + texto.replace("\\", "\\\\").replace('"', '\\"').replace("\n", " ") + '"'


def _md() -> MarkdownIt:
    return MarkdownIt("commonmark", {"html": False}).enable("table")


def _inline_typst(tokens) -> str:
    out = []
    for t in tokens:
        if t.type == "text":
            out.append(_s(t.content))
        elif t.type == "code_inline":
            out.append("#raw(" + _s(t.content)[1:] + ")")
        elif t.type == "softbreak":
            out.append(" ")
        elif t.type == "hardbreak":
            out.append(" \\\n")
        elif t.type == "strong_open":
            out.append("#strong[")
        elif t.type == "em_open":
            out.append("#emph[")
        elif t.type in ("strong_close", "em_close", "link_close"):
            out.append("]")
        elif t.type == "link_open":
            out.append("#link(" + _s(t.attrGet("href") or "")[1:] + ")[")
        elif t.type == "image":
            out.append(_s(f"[imagem: {t.content}]"))
        elif t.children:
            out.append(_inline_typst(t.children))
        elif t.content:
            out.append(_s(t.content))
    return "".join(out)


def markdown_para_typst(md: str) -> str:
    tokens = _md().parse(md)
    linhas: list[str] = []
    pilha_listas: list[str] = []
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            nivel = int(t.tag[1])
            linhas.append("=" * nivel + " " + _inline_typst(tokens[i + 1].children or []))
            i += 3
            continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            pilha_listas.append("-" if t.type == "bullet_list_open" else "+")
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            pilha_listas.pop()
            linhas.append("")
        elif t.type == "list_item_open":
            recuo = "  " * (len(pilha_listas) - 1)
            linhas.append(f"{recuo}{pilha_listas[-1]} ")
        elif t.type == "inline":
            texto = _inline_typst(t.children or [])
            if pilha_listas and linhas and linhas[-1].rstrip() in ("-", "+") or (linhas and re.fullmatch(r"\s*[-+] ", linhas[-1] or "")):
                linhas[-1] += texto
            else:
                linhas.append(texto)
        elif t.type == "paragraph_close" and not pilha_listas:
            linhas.append("")
        elif t.type in ("fence", "code_block"):
            corpo = t.content.rstrip("\n").replace("\\", "\\\\").replace('"', '\\"').replace("\n", "\\n")
            linhas.append(f'#raw(block: true, "{corpo}")')
            linhas.append("")
        elif t.type == "hr":
            linhas.append("#line(length: 100%)")
        elif t.type == "blockquote_open":
            linhas.append("#quote(block: true)[")
        elif t.type == "blockquote_close":
            linhas.append("]")
        elif t.type == "table_open":
            j, celulas, colunas = i + 1, [], 0
            linha_atual = 0
            while tokens[j].type != "table_close":
                if tokens[j].type == "tr_open":
                    linha_atual = 0
                if tokens[j].type in ("th_open", "td_open"):
                    linha_atual += 1
                    colunas = max(colunas, linha_atual)
                if tokens[j].type == "inline":
                    celulas.append("[" + _inline_typst(tokens[j].children or []) + "]")
                j += 1
            linhas.append(f"#table(columns: {max(colunas, 1)}, " + ", ".join(celulas) + ")")
            linhas.append("")
            i = j + 1
            continue
        i += 1
    return "\n".join(linhas)


def _preambulo(titulo: str, autor: str, idioma: str) -> str:
    lang, _, region = idioma.lower().partition("-")
    return (f"#set document(title: {_s(titulo)[1:]}, author: {_s(autor)[1:]})\n"
            f'#set text(font: "{FAMILIA}", size: 11pt, lang: "{lang or "pt"}"'
            + (f', region: "{region}"' if region else "") + ")\n"
            '#set page(paper: "a4", margin: 2.2cm, numbering: "1")\n'
            "#set par(justify: true)\n")


def _typst(fonte_typ: str, out: pathlib.Path, padroes: list[str] | None = None) -> None:
    exe = BF.qual("typst")
    if not exe:
        raise PapiroErro("E_SEM_SUPORTE", "Typst ausente")
    with tempfile.TemporaryDirectory() as t:
        typ = pathlib.Path(t) / "doc.typ"
        typ.write_text(fonte_typ, encoding="utf-8")
        cmd = [exe, "compile", "--font-path", str(FONTS), "--ignore-system-fonts"]
        if padroes:
            cmd += ["--pdf-standard", ",".join(padroes)]
        r = BF.rodar(cmd + [str(typ), str(out)], timeout=config.jobs("timeout_motor_s"), checar=False)
        if r.returncode != 0:
            msg = r.stderr.decode("utf-8", "replace").strip().splitlines()
            raise PapiroErro("E_MOTOR", f"typst: {(msg[0] if msg else 'falhou')[:300]}")


def _story(html: str, out: pathlib.Path) -> None:
    css = (f"@font-face {{font-family: papiro; src: url(LiberationSans-Regular.ttf);}}"
           f"@font-face {{font-family: papiro; src: url(LiberationSans-Bold.ttf); font-weight: bold;}}"
           f"@font-face {{font-family: papiro; src: url(LiberationSans-Italic.ttf); font-style: italic;}}"
           "* {font-family: papiro; font-size: 11pt; line-height: 1.35;} h1 {font-size: 20pt;} h2 {font-size: 16pt;}"
           "h3 {font-size: 13pt;} table {border-collapse: collapse;} td, th {border: 0.5pt solid #555; padding: 3pt;}")
    story = fitz.Story(html=html, user_css=css, archive=fitz.Archive(str(FONTS)))
    mediabox = fitz.paper_rect("a4")
    area = mediabox + (62, 62, -62, -62)
    writer = fitz.DocumentWriter(str(out))
    mais = 1
    while mais:
        dev = writer.begin_page(mediabox)
        mais, _ = story.place(area)
        story.draw(dev)
        writer.end_page()
    writer.close()


def markdown_para_pdf(md: str, out: pathlib.Path, titulo: str = "", autor: str = "PAPIRO",
                      idioma: str = "pt-BR", motor: str = "auto", padroes: list[str] | None = None) -> dict:
    """RF-301. motor: auto|typst|story. Grava titulo, autor, idioma e produtor sincronizados em Info e XMP."""
    if not md.strip():
        raise PapiroErro("E_ENTRADA", "markdown vazio")
    titulo = titulo or next((ln.lstrip("# ").strip() for ln in md.splitlines() if ln.startswith("# ")), "Documento")
    if not re.search(r"^# ", md, re.M):
        md = f"# {titulo}\n\n{md}"
    usar_typst = motor == "typst" or (motor == "auto" and BF.qual("typst"))
    fallback_from = None
    with tempfile.TemporaryDirectory() as t:
        bruto = pathlib.Path(t) / "bruto.pdf"
        via = "typst"
        try:
            if not usar_typst:
                raise PapiroErro("E_SEM_SUPORTE", "typst nao selecionado")
            _typst(_preambulo(titulo, autor, idioma) + markdown_para_typst(md), bruto, padroes)
        except PapiroErro as e:
            if motor == "typst" or padroes:
                raise
            if usar_typst:
                fallback_from = f"typst ({e.codigo})"
            via = "pymupdf-story"
            _story(_md().render(md), bruto)
        produtor = f"PAPIRO ({via} {ENG.versao('typst' if via == 'typst' else 'pymupdf') or ''})".strip()
        if padroes:
            with pikepdf.open(bruto) as pdf:
                with pdf.open_metadata(set_pikepdf_as_editor=False) as xmp:
                    xmp["pdf:Producer"] = produtor
                pdf.docinfo["/Producer"] = produtor
                pdf.save(out)
        else:
            metadados(bruto, out, titulo=titulo, autor=autor, idioma=idioma, produtor=produtor)
    with fitz.open(out) as d:
        n = d.page_count
    return {"ok": True, "via": via, "paginas": n, "fallback_from": fallback_from}


# ---------------- RF-304 ----------------
_ESCAPE_MD = re.compile(r"([\\`*_{}\[\]()#+\-.!|<>~])")


def carregar_dados(dados: str | pathlib.Path) -> list[dict]:
    if isinstance(dados, pathlib.Path):
        suf = dados.suffix.lower()
        if suf == ".json":
            reg = json.loads(dados.read_text(encoding="utf-8"))
        elif suf == ".csv":
            texto = dados.read_text(encoding="utf-8-sig")
            dialeto = csv.Sniffer().sniff(texto.splitlines()[0], delimiters=",;\t")
            reg = list(csv.DictReader(io.StringIO(texto), dialect=dialeto))
        elif suf in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            ws = load_workbook(dados, read_only=True, data_only=True).active
            linhas = list(ws.iter_rows(values_only=True))
            cab = [str(c) for c in linhas[0]]
            reg = [{cab[i]: ("" if v is None else v) for i, v in enumerate(l)} for l in linhas[1:] if any(v is not None for v in l)]
        else:
            raise PapiroErro("E_ENTRADA", f"formato de dados nao suportado: {suf}")
    else:
        try:
            reg = json.loads(dados or "[]")
        except json.JSONDecodeError as e:
            raise PapiroErro("E_ENTRADA", f"dados_json invalido: {e.msg}")
    if isinstance(reg, dict):
        reg = [reg]
    if not isinstance(reg, list) or not all(isinstance(r, dict) for r in reg):
        raise PapiroErro("E_ENTRADA", "dados devem ser uma lista de objetos")
    return reg


def _valor(registro: dict, caminho: str):
    atual = registro
    for parte in caminho.split("."):
        if not isinstance(atual, dict) or parte not in atual:
            raise KeyError(caminho)
        atual = atual[parte]
    return atual


def preencher_modelo(modelo: str, registro: dict) -> str:
    faltando = []

    def troca(m):
        try:
            return _ESCAPE_MD.sub(r"\\\1", str(_valor(registro, m.group(1).strip())))
        except KeyError:
            faltando.append(m.group(1).strip())
            return ""
    texto = re.sub(r"\{\{\s*([\w.]+)\s*\}\}", troca, modelo)
    if faltando:
        raise PapiroErro("E_ENTRADA", f"campos ausentes no registro: {sorted(set(faltando))}")
    return texto


def mala_direta(modelo_md: str, registros: list[dict], out_dir: pathlib.Path, modo: str = "consolidado",
                titulo: str = "Mala direta", nome_arquivo: str = "lote.pdf", campo_nome: str = "") -> dict:
    if not registros:
        raise PapiroErro("E_ENTRADA", "nenhum registro")
    if modo not in ("consolidado", "um_por_registro"):
        raise PapiroErro("E_ENTRADA", "modo deve ser consolidado ou um_por_registro")
    textos = [preencher_modelo(modelo_md, r) for r in registros]
    from ..caminhos import nome_seguro, saida_livre
    saidas = []
    if modo == "consolidado":
        out = saida_livre(out_dir, nome_arquivo)
        if BF.qual("typst"):
            corpo = "\n#pagebreak()\n".join(markdown_para_typst(t) for t in textos)
            with tempfile.TemporaryDirectory() as t:
                bruto = pathlib.Path(t) / "b.pdf"
                _typst(_preambulo(titulo, "PAPIRO", "pt-BR") + corpo, bruto)
                metadados(bruto, out, titulo=titulo, autor="PAPIRO", produtor=f"PAPIRO (typst {ENG.versao('typst')})")
            via = "typst"
        else:
            html = "".join(f'<div style="page-break-after: always">{_md().render(t)}</div>' for t in textos)
            with tempfile.TemporaryDirectory() as t:
                bruto = pathlib.Path(t) / "b.pdf"
                _story(html, bruto)
                metadados(bruto, out, titulo=titulo, autor="PAPIRO", produtor="PAPIRO (pymupdf-story)")
            via = "pymupdf-story"
        saidas.append(out)
    else:
        via = "pymupdf-story"
        for i, (t, r) in enumerate(zip(textos, registros), start=1):
            base = nome_seguro(str(r.get(campo_nome, "")) if campo_nome else "", f"registro-{i:05d}")
            out = saida_livre(out_dir, f"{base}.pdf")
            with tempfile.TemporaryDirectory() as tmp:
                bruto = pathlib.Path(tmp) / "b.pdf"
                _story(_md().render(t), bruto)
                metadados(bruto, out, titulo=f"{titulo} {i}", autor="PAPIRO", produtor="PAPIRO (pymupdf-story)")
            saidas.append(out)
    return {"via": via, "registros": len(registros), "saidas": saidas}


# ---------------- RF-303 ----------------
def office_para_pdf(entrada: pathlib.Path, out_dir: pathlib.Path, perfil_dir: pathlib.Path) -> dict:
    soffice = BF.qual("soffice") or BF.qual("soffice.com")
    if not soffice:
        raise PapiroErro("E_SEM_SUPORTE", "LibreOffice (soffice) nao instalado")
    with tempfile.TemporaryDirectory() as t:
        BF.rodar([soffice, f"-env:UserInstallation={perfil_dir.resolve().as_uri()}", "--headless", "--norestore",
                  "--convert-to", "pdf", "--outdir", t, str(entrada)], timeout=config.jobs("timeout_motor_s"))
        gerado = pathlib.Path(t) / (entrada.stem + ".pdf")
        if not gerado.exists():
            raise PapiroErro("E_MOTOR", "LibreOffice nao gerou o PDF")
        from ..caminhos import saida_livre
        destino = saida_livre(out_dir, entrada.stem.split("_", 1)[-1] + ".pdf")
        shutil.copyfile(gerado, destino)
    return {"saida": destino}


# ---------------- RF-309 ----------------
def _retificar(img_path: pathlib.Path, destino: pathlib.Path) -> bool:
    """Correcao de perspectiva: maior quadrilatero da foto vira retangulo. False se nao achar documento."""
    import cv2
    import numpy as np
    img = cv2.imread(str(img_path))
    if img is None:
        raise PapiroErro("E_ENTRADA", "imagem ilegivel")
    cinza = cv2.GaussianBlur(cv2.cvtColor(img, cv2.COLOR_BGR2GRAY), (5, 5), 0)
    bordas = cv2.Canny(cinza, 50, 150)
    bordas = cv2.dilate(bordas, np.ones((3, 3), np.uint8))
    contornos, _ = cv2.findContours(bordas, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    area_img = img.shape[0] * img.shape[1]
    for c in sorted(contornos, key=cv2.contourArea, reverse=True)[:5]:
        aprox = cv2.approxPolyDP(c, 0.02 * cv2.arcLength(c, True), True)
        if len(aprox) == 4 and cv2.contourArea(aprox) > 0.2 * area_img:
            pts = aprox.reshape(4, 2).astype("float32")
            s, d = pts.sum(1), np.diff(pts, axis=1).ravel()
            ordem = np.array([pts[s.argmin()], pts[d.argmin()], pts[s.argmax()], pts[d.argmax()]], dtype="float32")
            (tl, tr, br, bl) = ordem
            w = int(max(np.linalg.norm(br - bl), np.linalg.norm(tr - tl)))
            h = int(max(np.linalg.norm(tr - br), np.linalg.norm(tl - bl)))
            m = cv2.getPerspectiveTransform(ordem, np.array([[0, 0], [w - 1, 0], [w - 1, h - 1], [0, h - 1]], dtype="float32"))
            cv2.imwrite(str(destino), cv2.warpPerspective(img, m, (w, h)), [cv2.IMWRITE_JPEG_QUALITY, 92])
            return True
    return False


def imagens_para_pdf(imagens: list[pathlib.Path], out: pathlib.Path, retificar: bool = False) -> dict:
    import img2pdf
    avisos = []
    with tempfile.TemporaryDirectory() as t:
        usar = []
        for i, im in enumerate(imagens):
            if retificar:
                dst = pathlib.Path(t) / f"r{i}.jpg"
                if _retificar(im, dst):
                    usar.append(dst)
                    continue
                avisos.append(f"imagem {i + 1}: contorno do documento nao encontrado; usada sem retificar")
            usar.append(im)
        try:
            out.write_bytes(img2pdf.convert([str(p) for p in usar]))
            via = "opencv+img2pdf" if retificar else "img2pdf"
        except Exception:
            with fitz.open() as doc:
                for p in usar:
                    pix = fitz.Pixmap(str(p))
                    page = doc.new_page(width=pix.width, height=pix.height)
                    page.insert_image(page.rect, filename=str(p))
                doc.save(out)
            via = "pymupdf"
    with fitz.open(out) as d:
        n = d.page_count
    return {"via": via, "paginas": n, "avisos": avisos}


# ---------------- RF-310 ----------------
def decodificar_qr(pdf: pathlib.Path, pagina: int = 1) -> list[str]:
    import cv2
    import numpy as np
    with fitz.open(pdf) as d:
        pix = d[pagina - 1].get_pixmap(dpi=150)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    ok, textos, _pts, _ = cv2.QRCodeDetector().detectAndDecodeMulti(img[:, :, :3].copy())
    return [t for t in (textos or []) if t] if ok else []


def qr_pdf(conteudo: str, out: pathlib.Path, legenda: bool = True) -> dict:
    import segno
    if not conteudo:
        raise PapiroErro("E_ENTRADA", "conteudo do QR vazio")
    with tempfile.TemporaryDirectory() as t:
        png = pathlib.Path(t) / "qr.png"
        segno.make(conteudo, error="m").save(png, scale=10, border=4)
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_image(fitz.Rect(72, 72, 300, 300), filename=str(png))
            if legenda:
                page.insert_font(fontname="papiro", fontfile=str(FONTS / "LiberationSans-Regular.ttf"))
                page.insert_textbox(fitz.Rect(72, 310, 523, 420), conteudo[:300], fontname="papiro", fontsize=9)
            bruto = pathlib.Path(t) / "qr.pdf"
            doc.save(bruto)
        metadados(bruto, out, titulo="Codigo QR", autor="PAPIRO", produtor="PAPIRO (segno)")
    lidos = decodificar_qr(out)
    if conteudo not in lidos:
        raise PapiroErro("E_CONFORMIDADE", "QR gerado nao foi decodificado de volta com o mesmo conteudo")
    return {"paginas": 1, "decodificado": True}


# ---------------- RF-302 / RF-308 ----------------
def _chrome() -> str | None:
    for nome in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "chrome"):
        p = BF.qual(nome)
        if p:
            return p
    return None


def navegador_para_pdf(url: str, out: pathlib.Path, titulo: str = "", autor: str = "PAPIRO", idioma: str = "pt-BR") -> dict:
    """Chromium headless: PDF marcado (tags) e com marcadores a partir dos titulos."""
    exe = _chrome()
    if not exe:
        raise PapiroErro("E_SEM_SUPORTE", "Chrome/Chromium nao instalado")
    with tempfile.TemporaryDirectory() as t:
        bruto = pathlib.Path(t) / "b.pdf"
        BF.rodar([exe, "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
                  f"--user-data-dir={pathlib.Path(t) / 'perfil'}", "--no-pdf-header-footer",
                  "--generate-pdf-document-outline", f"--print-to-pdf={bruto}", url], timeout=180)
        if not bruto.exists() or bruto.stat().st_size == 0:
            raise PapiroErro("E_MOTOR", "Chromium nao gerou o PDF")
        with pikepdf.open(bruto) as pdf:
            titulo = titulo or str(pdf.docinfo.get("/Title", "") or "Documento")
        metadados(bruto, out, titulo=titulo, autor=autor, idioma=idioma,
                  produtor="PAPIRO (chromium headless)")
    with fitz.open(out) as d:
        return {"via": "chromium", "paginas": d.page_count}


def html_para_pdf(html: str, out: pathlib.Path, titulo: str = "", autor: str = "PAPIRO", idioma: str = "pt-BR") -> dict:
    if not html.strip():
        raise PapiroErro("E_ENTRADA", "html vazio")
    with tempfile.TemporaryDirectory() as t:
        arq = pathlib.Path(t) / "doc.html"
        arq.write_text(html, encoding="utf-8")
        return navegador_para_pdf(arq.as_uri(), out, titulo, autor, idioma)


# ---------------- RF-305 / RF-405 (diagramas vetoriais) ----------------
def diagrama_dot(fonte_dot: str, out: pathlib.Path, titulo: str = "Diagrama") -> dict:
    exe = BF.qual("dot")
    if not exe:
        raise PapiroErro("E_SEM_SUPORTE", "Graphviz (dot) nao instalado")
    with tempfile.TemporaryDirectory() as t:
        bruto = pathlib.Path(t) / "d.pdf"
        r = BF.rodar([exe, "-Tpdf", "-o", str(bruto)], timeout=120, checar=False, entrada=fonte_dot.encode("utf-8"))
        if r.returncode != 0 or not bruto.exists():
            msg = r.stderr.decode("utf-8", "replace").strip().splitlines()
            raise PapiroErro("E_ENTRADA", f"DOT invalido: {(msg[0] if msg else '')[:200]}")
        metadados(bruto, out, titulo=titulo, autor="PAPIRO", produtor=f"PAPIRO (graphviz)")
    return {"via": "graphviz", "paginas": 1, "vetorial": True}
